import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from database import SessionDB, SetDB
from sqlalchemy.orm import Session

XLSX_FILENAME = "mahmoud_gym_tracker.xlsx"

def generate_xlsx(db: Session, filepath: str = XLSX_FILENAME):
    wb = openpyxl.Workbook()
    
    # 1. Summary Sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"
    
    headers_summary = ["Session ID", "Date", "Time", "Session Type", "Bodyweight (kg)", "Total Volume", "Total Sets", "Duration (s)"]
    ws_summary.append(headers_summary)
    
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    
    for col_num, _ in enumerate(headers_summary, 1):
        cell = ws_summary.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
    
    sessions = db.query(SessionDB).order_by(SessionDB.date.asc(), SessionDB.time.asc()).all()
    for s in sessions:
        ws_summary.append([
            s.id, s.date, s.time, s.session_type, s.bodyweight, s.total_volume, s.total_sets, s.duration
        ])
    
    # 2. All Sets Sheet
    ws_sets = wb.create_sheet(title="All Sets")
    headers_sets = ["Session ID", "Date", "Session Type", "Exercise", "Set Num", "Weight (kg)", "Reps", "Volume"]
    ws_sets.append(headers_sets)
    for col_num, _ in enumerate(headers_sets, 1):
        cell = ws_sets.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        
    sets_data = db.query(SetDB, SessionDB).join(SessionDB).order_by(SessionDB.date.asc(), SessionDB.time.asc(), SetDB.id.asc()).all()
    for st, sess in sets_data:
        ws_sets.append([
            sess.id, sess.date, sess.session_type, st.exercise_name, st.set_num, st.weight_kg, st.reps, st.volume
        ])
        
    # 3. Per Session Type Sheet
    ws_type = wb.create_sheet(title="Per Session Type")
    headers_type = ["Session Type", "Count", "Avg Volume", "Avg Sets", "Avg Duration (s)"]
    ws_type.append(headers_type)
    for col_num, _ in enumerate(headers_type, 1):
        cell = ws_type.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
    
    # Aggregate data
    types = {}
    for s in sessions:
        if s.session_type not in types:
            types[s.session_type] = {"count": 0, "total_vol": 0, "total_sets": 0, "total_dur": 0}
        types[s.session_type]["count"] += 1
        types[s.session_type]["total_vol"] += s.total_volume
        types[s.session_type]["total_sets"] += s.total_sets
        types[s.session_type]["total_dur"] += s.duration
        
    for t, data in types.items():
        count = data["count"]
        ws_type.append([
            t, 
            count, 
            round(data["total_vol"] / count, 2), 
            round(data["total_sets"] / count, 2), 
            round(data["total_dur"] / count, 2)
        ])
    
    wb.save(filepath)
    return filepath
